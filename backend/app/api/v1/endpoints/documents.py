"""
Document API Endpoints
"""
import os
import uuid
import logging
from datetime import datetime
from typing import List, Optional
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Query
from sqlalchemy.orm import Session
from app.db.session import get_db
from app.models.document import Document, ProcessingStatus
from app.schemas.document import (
    DocumentResponse,
    DocumentCreate,
    DocumentUpdate,
    DocumentClassification,
    LayoutAnalysis,
    EntityExtraction,
    InformationExtraction,
    AISummary,
    ConfidenceScores,
    DocumentSearch,
    ExportRequest,
    ExportResponse,
    DashboardStats,
)
from app.services.ai import (
    ocr_service,
    classification_service,
    layout_service,
    ner_service,
    extraction_service,
    summary_service,
)
from app.core.config import settings

logger = logging.getLogger(__name__)

router = APIRouter()


@router.post("/upload", response_model=DocumentResponse)
async def upload_document(
    file: UploadFile = File(...),
    db: Session = Depends(get_db)
):
    """
    Upload a document for processing
    Supports: PDF, JPG, JPEG, PNG, TIFF
    """
    # Validate file type
    if file.content_type not in [
        "application/pdf",
        "image/jpeg",
        "image/jpg",
        "image/png",
        "image/tiff",
        "image/bmp",
        "image/webp"
    ]:
        raise HTTPException(status_code=400, detail="Invalid file type")
    
    # Generate unique filename
    file_extension = file.filename.split(".")[-1].lower()
    unique_filename = f"{uuid.uuid4()}.{file_extension}"
    file_path = os.path.join(settings.UPLOAD_DIR, unique_filename)
    
    # Save file
    try:
        with open(file_path, "wb") as buffer:
            content = await file.read()
            if len(content) > settings.MAX_UPLOAD_SIZE:
                raise HTTPException(status_code=413, detail="File size exceeds limit")
            buffer.write(content)
    except Exception as e:
        logger.error(f"Failed to save file: {e}")
        raise HTTPException(status_code=500, detail="Failed to save file")
    
    # Create document record
    document = Document(
        filename=unique_filename,
        original_filename=file.filename,
        file_path=file_path,
        file_size=len(content),
        mime_type=file.content_type,
        status=ProcessingStatus.PENDING,
        created_at=datetime.utcnow()
    )
    
    db.add(document)
    db.commit()
    db.refresh(document)
    
    # Trigger async processing (for now, process synchronously)
    try:
        await process_document(document, db)
    except Exception as e:
        logger.error(f"Document processing failed: {e}")
        document.status = ProcessingStatus.FAILED
        db.commit()
    
    db.refresh(document)
    return document


async def process_document(document: Document, db: Session):
    """
    Process document through the AI pipeline
    """
    document.status = ProcessingStatus.PROCESSING
    document.processing_started_at = datetime.utcnow()
    db.commit()
    
    start_time = datetime.utcnow()
    
    try:
        # Step 1: OCR
        logger.info(f"Processing OCR for document {document.id}")
        if document.mime_type == "application/pdf":
            ocr_result = ocr_service.extract_text_from_pdf(document.file_path)
        else:
            ocr_result = ocr_service.extract_text(document.file_path)
        
        document.ocr_text = ocr_result["text"]
        document.ocr_confidence = ocr_result["confidence"]
        document.language_detected = ocr_result["language_detected"]
        document.ocr_confidence_score = ocr_result["confidence"]
        
        # Step 2: Classification
        logger.info(f"Classifying document {document.id}")
        classification_result = classification_service.classify(
            document.ocr_text,
            document.file_path
        )
        
        document.document_type = classification_result["document_type"]
        document.classification_confidence = classification_result["confidence"]
        document.classification_probabilities = classification_result["probabilities"]
        document.classification_confidence_score = classification_result["confidence"]
        
        # Step 3: Layout Analysis
        logger.info(f"Analyzing layout for document {document.id}")
        layout_result = layout_service.analyze_layout(
            document.file_path,
            ocr_result["regions"]
        )
        
        document.layout_analysis = layout_result
        document.layout_confidence_score = 0.7  # Fixed for now
        
        # Step 4: NER
        logger.info(f"Extracting entities for document {document.id}")
        ner_result = ner_service.extract_entities(document.ocr_text)
        
        # Color-code entities
        color_coded_entities = ner_service.color_code_entities(
            document.ocr_text,
            ner_result["entities"]
        )
        
        document.entities = color_coded_entities
        document.entity_confidence_score = ner_result["confidence"]
        
        # Step 5: Information Extraction
        logger.info(f"Extracting information for document {document.id}")
        extraction_result = extraction_service.extract_information(
            document.ocr_text,
            document.document_type,
            document.entities
        )
        
        document.extracted_data = extraction_result
        document.extraction_confidence_score = extraction_result["confidence"]
        
        # Step 6: AI Summary
        logger.info(f"Generating summary for document {document.id}")
        summary_result = summary_service.generate_summary(
            document.ocr_text,
            document.document_type,
            extraction_result,
            document.entities
        )
        
        document.summary = summary_result["summary"]
        
        # Calculate overall confidence
        document.overall_confidence_score = (
            document.ocr_confidence_score * 0.25 +
            document.classification_confidence_score * 0.2 +
            document.layout_confidence_score * 0.15 +
            document.entity_confidence_score * 0.2 +
            document.extraction_confidence_score * 0.2
        )
        
        # Update status
        document.status = ProcessingStatus.COMPLETED
        document.processing_completed_at = datetime.utcnow()
        document.processing_duration = (
            document.processing_completed_at - document.processing_started_at
        ).total_seconds()
        
        db.commit()
        
        logger.info(f"Document {document.id} processed successfully")
        
    except Exception as e:
        logger.error(f"Error processing document {document.id}: {e}")
        document.status = ProcessingStatus.FAILED
        document.processing_completed_at = datetime.utcnow()
        db.commit()
        raise


@router.get("/documents", response_model=List[DocumentResponse])
def get_documents(
    skip: int = Query(0, ge=0),
    limit: int = Query(100, ge=1, le=100),
    db: Session = Depends(get_db)
):
    """Get all documents with pagination"""
    documents = db.query(Document).order_by(Document.created_at.desc()).offset(skip).limit(limit).all()
    return documents


@router.get("/document/{document_id}", response_model=DocumentResponse)
def get_document(document_id: int, db: Session = Depends(get_db)):
    """Get a specific document by ID"""
    document = db.query(Document).filter(Document.id == document_id).first()
    if not document:
        raise HTTPException(status_code=404, detail="Document not found")
    return document


@router.delete("/document/{document_id}")
def delete_document(document_id: int, db: Session = Depends(get_db)):
    """Delete a document"""
    document = db.query(Document).filter(Document.id == document_id).first()
    if not document:
        raise HTTPException(status_code=404, detail="Document not found")
    
    # Delete file
    if os.path.exists(document.file_path):
        os.remove(document.file_path)
    
    db.delete(document)
    db.commit()
    
    return {"message": "Document deleted successfully"}


@router.post("/classify", response_model=DocumentClassification)
async def classify_document(
    file: UploadFile = File(...),
):
    """Classify a document without saving it"""
    # Save temporarily
    temp_filename = f"temp_{uuid.uuid4()}"
    temp_path = os.path.join(settings.UPLOAD_DIR, temp_filename)
    
    try:
        with open(temp_path, "wb") as buffer:
            content = await file.read()
            buffer.write(content)
        
        # Get OCR text
        if file.content_type == "application/pdf":
            ocr_result = ocr_service.extract_text_from_pdf(temp_path)
        else:
            ocr_result = ocr_service.extract_text(temp_path)
        
        # Classify
        result = classification_service.classify(ocr_result["text"], temp_path)
        
        return result
        
    finally:
        if os.path.exists(temp_path):
            os.remove(temp_path)


@router.post("/layout", response_model=LayoutAnalysis)
async def analyze_layout(
    file: UploadFile = File(...),
):
    """Analyze document layout without saving it"""
    temp_filename = f"temp_{uuid.uuid4()}"
    temp_path = os.path.join(settings.UPLOAD_DIR, temp_filename)
    
    try:
        with open(temp_path, "wb") as buffer:
            content = await file.read()
            buffer.write(content)
        
        # Get OCR
        ocr_result = ocr_service.extract_text(temp_path)
        
        # Analyze layout
        result = layout_service.analyze_layout(temp_path, ocr_result["regions"])
        
        return result
        
    finally:
        if os.path.exists(temp_path):
            os.remove(temp_path)


@router.post("/ocr")
async def extract_ocr(
    file: UploadFile = File(...),
):
    """Extract text from document without saving it"""
    temp_filename = f"temp_{uuid.uuid4()}"
    temp_path = os.path.join(settings.UPLOAD_DIR, temp_filename)
    
    try:
        with open(temp_path, "wb") as buffer:
            content = await file.read()
            buffer.write(content)
        
        # Extract text
        if file.content_type == "application/pdf":
            result = ocr_service.extract_text_from_pdf(temp_path)
        else:
            result = ocr_service.extract_text(temp_path)
        
        return result
        
    finally:
        if os.path.exists(temp_path):
            os.remove(temp_path)


@router.post("/extract", response_model=InformationExtraction)
async def extract_information(
    file: UploadFile = File(...),
):
    """Extract information from document without saving it"""
    temp_filename = f"temp_{uuid.uuid4()}"
    temp_path = os.path.join(settings.UPLOAD_DIR, temp_filename)
    
    try:
        with open(temp_path, "wb") as buffer:
            content = await file.read()
            buffer.write(content)
        
        # Get OCR
        ocr_result = ocr_service.extract_text(temp_path)
        
        # Classify
        classification_result = classification_service.classify(ocr_result["text"], temp_path)
        
        # Get entities
        ner_result = ner_service.extract_entities(ocr_result["text"])
        
        # Extract information
        result = extraction_service.extract_information(
            ocr_result["text"],
            classification_result["document_type"],
            ner_result["entities"]
        )
        
        return result
        
    finally:
        if os.path.exists(temp_path):
            os.remove(temp_path)


@router.post("/entities", response_model=EntityExtraction)
async def extract_entities(
    file: UploadFile = File(...),
):
    """Extract named entities from document without saving it"""
    temp_filename = f"temp_{uuid.uuid4()}"
    temp_path = os.path.join(settings.UPLOAD_DIR, temp_filename)
    
    try:
        with open(temp_path, "wb") as buffer:
            content = await file.read()
            buffer.write(content)
        
        # Get OCR
        ocr_result = ocr_service.extract_text(temp_path)
        
        # Extract entities
        result = ner_service.extract_entities(ocr_result["text"])
        
        return result
        
    finally:
        if os.path.exists(temp_path):
            os.remove(temp_path)


@router.post("/summary", response_model=AISummary)
async def generate_summary(
    file: UploadFile = File(...),
):
    """Generate AI summary for document without saving it"""
    temp_filename = f"temp_{uuid.uuid4()}"
    temp_path = os.path.join(settings.UPLOAD_DIR, temp_filename)
    
    try:
        with open(temp_path, "wb") as buffer:
            content = await file.read()
            buffer.write(content)
        
        # Get OCR
        ocr_result = ocr_service.extract_text(temp_path)
        
        # Classify
        classification_result = classification_service.classify(ocr_result["text"], temp_path)
        
        # Get entities
        ner_result = ner_service.extract_entities(ocr_result["text"])
        
        # Extract information
        extraction_result = extraction_service.extract_information(
            ocr_result["text"],
            classification_result["document_type"],
            ner_result["entities"]
        )
        
        # Generate summary
        result = summary_service.generate_summary(
            ocr_result["text"],
            classification_result["document_type"],
            extraction_result,
            ner_result["entities"]
        )
        
        return result
        
    finally:
        if os.path.exists(temp_path):
            os.remove(temp_path)


@router.get("/dashboard/stats", response_model=DashboardStats)
def get_dashboard_stats(db: Session = Depends(get_db)):
    """Get dashboard statistics"""
    total_documents = db.query(Document).count()
    processing_documents = db.query(Document).filter(
        Document.status == ProcessingStatus.PROCESSING
    ).count()
    completed_documents = db.query(Document).filter(
        Document.status == ProcessingStatus.COMPLETED
    ).count()
    failed_documents = db.query(Document).filter(
        Document.status == ProcessingStatus.FAILED
    ).count()
    
    # Calculate average confidence
    completed_docs = db.query(Document).filter(
        Document.status == ProcessingStatus.COMPLETED
    ).all()
    
    if completed_docs:
        avg_confidence = sum(doc.overall_confidence_score for doc in completed_docs) / len(completed_docs)
        avg_processing_time = sum(
            doc.processing_duration for doc in completed_docs 
            if doc.processing_duration
        ) / len([doc for doc in completed_docs if doc.processing_duration])
    else:
        avg_confidence = 0.0
        avg_processing_time = 0.0
    
    # Document type distribution
    type_distribution = {}
    for doc in completed_docs:
        doc_type = doc.document_type.value if doc.document_type else "unknown"
        type_distribution[doc_type] = type_distribution.get(doc_type, 0) + 1
    
    # Recent uploads
    recent_uploads = db.query(Document).order_by(
        Document.created_at.desc()
    ).limit(10).all()
    
    return DashboardStats(
        total_documents=total_documents,
        processing_documents=processing_documents,
        completed_documents=completed_documents,
        failed_documents=failed_documents,
        average_confidence=avg_confidence,
        average_processing_time=avg_processing_time,
        document_type_distribution=type_distribution,
        recent_uploads=recent_uploads
    )


@router.post("/export/json")
async def export_json(
    document_id: int,
    db: Session = Depends(get_db)
):
    """Export document as JSON"""
    document = db.query(Document).filter(Document.id == document_id).first()
    if not document:
        raise HTTPException(status_code=404, detail="Document not found")
    
    export_data = {
        "document_id": document.id,
        "filename": document.filename,
        "original_filename": document.original_filename,
        "document_type": document.document_type.value if document.document_type else None,
        "classification_confidence": document.classification_confidence,
        "ocr_text": document.ocr_text,
        "ocr_confidence": document.ocr_confidence,
        "language_detected": document.language_detected,
        "layout_analysis": document.layout_analysis,
        "entities": document.entities,
        "extracted_data": document.extracted_data,
        "summary": document.summary,
        "confidence_scores": {
            "ocr": document.ocr_confidence_score,
            "classification": document.classification_confidence_score,
            "layout": document.layout_confidence_score,
            "entity": document.entity_confidence_score,
            "extraction": document.extraction_confidence_score,
            "overall": document.overall_confidence_score
        },
        "created_at": document.created_at.isoformat() if document.created_at else None,
        "processing_duration": document.processing_duration
    }
    
    return export_data


@router.post("/export/csv")
async def export_csv(
    document_id: int,
    db: Session = Depends(get_db)
):
    """Export document as CSV"""
    document = db.query(Document).filter(Document.id == document_id).first()
    if not document:
        raise HTTPException(status_code=404, detail="Document not found")
    
    # Simple CSV export of extracted data
    import csv
    import io
    
    output = io.StringIO()
    writer = csv.writer(output)
    
    # Write header
    writer.writerow(["Field", "Value"])
    
    # Write basic info
    writer.writerow(["Document ID", document.id])
    writer.writerow(["Filename", document.original_filename])
    writer.writerow(["Document Type", document.document_type.value if document.document_type else ""])
    writer.writerow(["OCR Confidence", document.ocr_confidence])
    writer.writerow(["Overall Confidence", document.overall_confidence_score])
    
    # Write extracted data
    if document.extracted_data:
        extracted = document.extracted_data.get("extracted_data", {})
        for key, value in extracted.items():
            if value:
                writer.writerow([key, str(value)])
    
    # Write summary
    if document.summary:
        writer.writerow(["Summary", document.summary])
    
    csv_content = output.getvalue()
    output.close()
    
    return {"csv_content": csv_content}


def _safe_format_confidence(val):
    if val is None:
        return "0.0%"
    if isinstance(val, str):
        s = val.strip()
        if s.endswith("%"):
            return s
        try:
            val = float(s)
        except ValueError:
            return f"{s}%"
    if isinstance(val, (int, float)):
        v = val * 100.0 if val <= 1.0 else val
        return f"{v:.1f}%"
    return f"{val}%"


@router.post("/export/txt")
async def export_txt(
    document_id: int,
    db: Session = Depends(get_db)
):
    """Export document as TXT (legacy JSON response containing txt content)"""
    document = db.query(Document).filter(Document.id == document_id).first()
    if not document:
        raise HTTPException(status_code=404, detail="Document not found")
    
    ocr_conf_str = _safe_format_confidence(document.ocr_confidence)
    overall_conf_str = _safe_format_confidence(document.overall_confidence_score)

    txt_content = f"""
Document ID: {document.id}
Filename: {document.original_filename}
Document Type: {document.document_type.value if document.document_type else 'Unknown'}
OCR Confidence: {ocr_conf_str}
Overall Confidence: {overall_conf_str}

--- OCR Text ---
{document.ocr_text}

--- Summary ---
{document.summary}

--- Extracted Data ---
{document.extracted_data}
"""
    
    return {"txt_content": txt_content}


# New endpoints that return actual downloadable files (binary responses)
from fastapi.responses import StreamingResponse
import io
import json
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib import colors
from docx import Document as DocxDocument
from docx.shared import Pt
from openpyxl import Workbook


@router.get("/document/{document_id}/export/{fmt}")
async def export_document_file(document_id: int, fmt: str, db: Session = Depends(get_db)):
    """Export document as a downloadable file: txt, json, csv, pdf, docx, xlsx"""
    document = db.query(Document).filter(Document.id == document_id).first()
    if not document:
        raise HTTPException(status_code=404, detail="Document not found")

    fmt = fmt.lower()
    safe_name = (document.original_filename or f"document_{document.id}").replace(' ', '_')

    # Prepare common data
    export_obj = {
        "document_id": document.id,
        "filename": document.original_filename,
        "document_type": document.document_type.value if document.document_type else None,
        "summary": document.summary,
        "raw_text": document.ocr_text,
        "formatted_text": None,
        "entities": document.entities,
        "tables": document.layout_analysis.get('tables') if isinstance(document.layout_analysis, dict) and 'tables' in document.layout_analysis else [],
        "layout": document.layout_analysis,
        "structured_data": document.extracted_data,
    }

    if fmt in ["txt", "text"]:
        content = io.BytesIO()
        text = []
        text.append(f"Document ID: {document.id}")
        text.append(f"Filename: {document.original_filename}")
        text.append(f"Document Type: {document.document_type.value if document.document_type else 'Unknown'}")
        text.append("")
        text.append("--- OCR Text ---")
        text.append(document.ocr_text or "")
        text.append("")
        text.append("--- Summary ---")
        text.append(document.summary or "")
        text.append("")
        text.append("--- Extracted Data ---")
        text.append(json.dumps(document.extracted_data or {}, indent=2, default=str))
        final = "\n".join(text)
        content.write(final.encode("utf-8"))
        content.seek(0)
        headers = {"Content-Disposition": f"attachment; filename=AI_Report_{safe_name}.txt"}
        return StreamingResponse(content, media_type="text/plain; charset=utf-8", headers=headers)

    if fmt == "json":
        content = io.BytesIO()
        content.write(json.dumps(export_obj, indent=2, default=str).encode("utf-8"))
        content.seek(0)
        headers = {"Content-Disposition": f"attachment; filename=AI_Report_{safe_name}.json"}
        return StreamingResponse(content, media_type="application/json", headers=headers)

    if fmt == "csv":
        import csv
        content = io.StringIO()
        writer = csv.writer(content)
        writer.writerow(["Field", "Value"])
        writer.writerow(["Document ID", document.id])
        writer.writerow(["Filename", document.original_filename])
        writer.writerow(["Document Type", document.document_type.value if document.document_type else ""]) 
        writer.writerow(["Summary", document.summary or ""]) 
        if document.extracted_data:
            extracted = document.extracted_data.get("extracted_data", {}) if isinstance(document.extracted_data, dict) else {}
            for key, value in extracted.items():
                writer.writerow([key, json.dumps(value, default=str)])
        csv_bytes = io.BytesIO(content.getvalue().encode("utf-8"))
        csv_bytes.seek(0)
        headers = {"Content-Disposition": f"attachment; filename=AI_Report_{safe_name}.csv"}
        return StreamingResponse(csv_bytes, media_type="text/csv", headers=headers)

    if fmt == "pdf":
        buf = io.BytesIO()
        doc_template = SimpleDocTemplate(buf, pagesize=A4, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
        styles = getSampleStyleSheet()
        normal = styles['Normal']
        normal.fontSize = 11
        normal.leading = 14
        story = []

        # Title
        title_style = styles['Title']
        story.append(Paragraph(f"Document Report - {document.original_filename}", title_style))
        story.append(Spacer(1, 12))

        # Document Info
        story.append(Paragraph("Document Information", styles['Heading2']))
        info_text = f"Document Type: {document.document_type.value if document.document_type else 'Unknown'}<br/>"
        info_text += f"Processed At: {document.processing_completed_at.isoformat() if document.processing_completed_at else 'N/A'}"
        story.append(Paragraph(info_text, normal))
        story.append(Spacer(1, 12))

        # Summary
        story.append(Paragraph("Summary", styles['Heading3']))
        story.append(Paragraph(document.summary or "No summary available", normal))
        story.append(Spacer(1, 12))

        # Entities
        story.append(Paragraph("Named Entities", styles['Heading3']))
        entities = document.entities or []
        if isinstance(entities, list) and entities:
            data = [["Label", "Text", "Confidence"]]
            for e in entities:
                label = e.get('label') if isinstance(e, dict) else ''
                text = e.get('text') if isinstance(e, dict) else str(e)
                confidence = e.get('confidence', '') if isinstance(e, dict) else ''
                data.append([label, text, str(confidence)])
            table = Table(data, colWidths=[100, 300, 80])
            table.setStyle(TableStyle([
                ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#f3f4f6')),
                ('GRID', (0,0), (-1,-1), 0.5, colors.black),
                ('VALIGN', (0,0), (-1,-1), 'TOP')
            ]))
            story.append(table)
        else:
            story.append(Paragraph('No entities found', normal))
        story.append(Spacer(1, 12))

        # Tables (from layout)
        story.append(Paragraph("Detected Tables", styles['Heading3']))
        tables = export_obj.get('tables') or []
        if isinstance(tables, list) and tables:
            for t in tables:
                # Expect t as list of rows
                try:
                    tbl = Table(t)
                    tbl.setStyle(TableStyle([
                        ('GRID', (0,0), (-1,-1), 0.5, colors.black),
                        ('VALIGN', (0,0), (-1,-1), 'TOP')
                    ]))
                    story.append(tbl)
                    story.append(Spacer(1,12))
                except Exception:
                    continue
        else:
            story.append(Paragraph('No tables detected', normal))
        story.append(Spacer(1, 12))

        # Formatted Document Content
        story.append(Paragraph("Formatted Document Content", styles['Heading3']))
        formatted_text = document.ocr_text or ''
        # Split into paragraphs by double newline
        for para in (formatted_text.split('\n\n') if formatted_text else []):
            story.append(Paragraph(para.replace('\n', '<br/>'), normal))
            story.append(Spacer(1,6))

        doc_template.build(story)
        buf.seek(0)
        headers = {"Content-Disposition": f"attachment; filename=AI_Report_{safe_name}.pdf"}
        return StreamingResponse(buf, media_type="application/pdf", headers=headers)

    if fmt == "docx":
        buf = io.BytesIO()
        docx = DocxDocument()
        style = docx.styles['Normal']
        font = style.font
        font.name = 'Calibri'
        font.size = Pt(11)

        docx.add_heading(f"Document Report - {document.original_filename}", level=1)
        docx.add_heading("Document Information", level=2)
        p = docx.add_paragraph()
        p.add_run(f"Document Type: {document.document_type.value if document.document_type else 'Unknown'}\n")
        p.add_run(f"Processed At: {document.processing_completed_at.isoformat() if document.processing_completed_at else 'N/A'}\n")

        docx.add_heading("Summary", level=2)
        docx.add_paragraph(document.summary or "No summary available")

        docx.add_heading("Named Entities", level=2)
        entities = document.entities or []
        if isinstance(entities, list) and entities:
            table = docx.add_table(rows=1, cols=3)
            hdr_cells = table.rows[0].cells
            hdr_cells[0].text = 'Label'
            hdr_cells[1].text = 'Text'
            hdr_cells[2].text = 'Confidence'
            for e in entities:
                row_cells = table.add_row().cells
                row_cells[0].text = str(e.get('label') if isinstance(e, dict) else '')
                row_cells[1].text = str(e.get('text') if isinstance(e, dict) else str(e))
                row_cells[2].text = str(e.get('confidence',''))
        else:
            docx.add_paragraph('No entities found')

        docx.add_heading('Formatted Document Content', level=2)
        for para in ( (document.ocr_text or '').split('\n\n') ):
            docx.add_paragraph(para)

        docx.save(buf)
        buf.seek(0)
        headers = {"Content-Disposition": f"attachment; filename=AI_Report_{safe_name}.docx"}
        return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document", headers=headers)

    if fmt in ["xlsx", "excel"]:
        buf = io.BytesIO()
        wb = Workbook()
        ws = wb.active
        ws.title = 'Summary'
        ws.append(['Document ID', document.id])
        ws.append(['Filename', document.original_filename])
        ws.append(['Document Type', document.document_type.value if document.document_type else ''])
        ws.append(['Summary', document.summary or ''])

        # Entities sheet
        ent_ws = wb.create_sheet('Entities')
        ent_ws.append(['Label', 'Text', 'Confidence'])
        for e in (document.entities or []):
            ent_ws.append([e.get('label','') if isinstance(e, dict) else '', e.get('text','') if isinstance(e, dict) else str(e), e.get('confidence','') if isinstance(e, dict) else ''])

        # Extracted fields
        ext_ws = wb.create_sheet('Extracted Fields')
        ext_ws.append(['Field', 'Value'])
        if document.extracted_data and isinstance(document.extracted_data, dict):
            extracted = document.extracted_data.get('extracted_data', document.extracted_data)
            if isinstance(extracted, dict):
                for k, v in extracted.items():
                    ext_ws.append([k, json.dumps(v, default=str)])

        wb.save(buf)
        buf.seek(0)
        headers = {"Content-Disposition": f"attachment; filename=AI_Report_{safe_name}.xlsx"}
        return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers=headers)

    raise HTTPException(status_code=400, detail="Unsupported export format")

